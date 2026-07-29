"""Geração server-side dos documentos do kit (PDF via docxtpl + LibreOffice).

Porta a lógica de `montarContexto()` do NovoKitView.vue para Python e
gera todos os PDFs do kit armazenando-os como registros DocumentoKit.
"""
from __future__ import annotations

import re
import subprocess
import tempfile
import unicodedata
from datetime import date
from decimal import ROUND_HALF_UP, Decimal
from io import BytesIO
from pathlib import Path

from django.core.files.base import ContentFile

from templates_app.models import Template
from templates_app.docx_jinja_normalizer import normalize_docx_jinja_runs
from templates_app.docx_cleaner import strip_blank_pages
from templates_app.docx_style_flattener import flatten_inherited_formatting
from common.jinja_env import build_env
from common.bold_markers import aplicar_marcadores_negrito, marcar_negrito
from .models import DocumentoKit, Kit

try:
    from docxtpl import DocxTemplate
except ImportError:
    DocxTemplate = None


# ---------------------------------------------------------------------------
# Constantes
# ---------------------------------------------------------------------------

MESES = [
    "janeiro", "fevereiro", "março", "abril", "maio", "junho",
    "julho", "agosto", "setembro", "outubro", "novembro", "dezembro",
]

TIPOS_ACAO_LABELS = {
    "cartao_credito_consignado": "Cartão de Crédito Consignado",
    "rmc": "RMC",
    "rcc": "RCC",
    "emprestimo_nao_autorizado": "Empréstimo Não Autorizado",
    "tarifa_bancaria": "Tarifa Bancária",
    "contribuicao_sindical_nao_autorizada": "Contribuição Sindical Não Autorizada",
    "seguro_nao_autorizado": "Seguro Não Autorizado",
}

RELACAO_MAP = {
    "pai_mae": "pai/mãe",
    "filho_a": "filho(a)",
    "irmao_a": "irmão/irmã",
    "conjuge": "cônjuge",
}

# Tipos de ação que exibem número de contrato na procuração
TIPOS_COM_CONTRATO = {
    "cartao_credito_consignado",
    "rmc",
    "rcc",
    "emprestimo_nao_autorizado",
}

# Keywords por tipo de kit (mesma lógica do KIT_TEMPLATE_DEFS_* do frontend)
KIT_TEMPLATE_DEFS: dict[str, list[dict]] = {
    "bancario": [
        {"key": "contrato", "keyword": "kit contrato"},
        {"key": "procuracao", "keyword": "kit procuracao"},
        {"key": "hipossuficiencia", "keyword": "kit hipossuficiencia"},
        {"key": "ciencia", "keyword": "kit ciencia"},
        {"key": "domicilio", "keyword": "kit domicilio"},
    ],
    "previdenciario": [
        {"key": "contrato", "keyword": "kit contrato previdenciario"},
        {"key": "procuracao", "keyword": "kit procuracao previdenciario"},
        {"key": "hipossuficiencia", "keyword": "kit hipossuficiencia previdenciario"},
        {"key": "domicilio", "keyword": "kit domicilio previdenciario"},
    ],
    "marketing": [
        {"key": "contrato", "keyword": "kit contrato marketing"},
        {"key": "procuracao", "keyword": "kit procuracao marketing"},
        {"key": "hipossuficiencia", "keyword": "kit hipossuficiencia marketing"},
        {"key": "ciencia", "keyword": "kit ciencia marketing"},
    ],
}


# ---------------------------------------------------------------------------
# Helpers de formatação
# ---------------------------------------------------------------------------

def _normalize(s: str) -> str:
    return unicodedata.normalize("NFD", s).encode("ascii", "ignore").decode("ascii").lower().strip()


def slug_nome_cliente(nome: str) -> str:
    """'José da Silva' → 'JOSE_DA_SILVA' (para nomes de arquivo)."""
    sem_acento = unicodedata.normalize("NFD", nome).encode("ascii", "ignore").decode("ascii")
    sem_acento = re.sub(r"[^a-zA-Z0-9\s]", "", sem_acento)
    return "_".join(sem_acento.upper().split())


def _fmt_cpf(v: str) -> str:
    d = re.sub(r"\D", "", str(v or ""))[:11]
    return f"{d[:3]}.{d[3:6]}.{d[6:9]}-{d[9:]}" if len(d) == 11 else v or ""


def _fmt_cep(v: str) -> str:
    d = re.sub(r"\D", "", str(v or ""))[:8]
    return f"{d[:5]}-{d[5:]}" if len(d) == 8 else v or ""


def _juntar(items: list[str]) -> str:
    items = [i for i in items if i]
    if not items:
        return ""
    if len(items) == 1:
        return items[0]
    return ", ".join(items[:-1]) + " e " + items[-1]


def _formatar_relacao(tipo: str, outro: str) -> str:
    if tipo == "outro":
        return (outro or "").strip()
    return RELACAO_MAP.get(tipo, "")


def _qualificar_advogado(adv: dict) -> str:
    genero = adv.get("genero", "masculino")
    inscrito = "inscrita" if genero == "feminino" else "inscrito"
    texto = (
        f"{adv.get('nome_completo', '')}, "
        f"{adv.get('nacionalidade', '')}, "
        f"{adv.get('estado_civil', '')}, "
        f"advogado, {inscrito} na {adv.get('numero_oab', '')}"
    )
    escritorio_nome = adv.get("escritorio_nome", "")
    escritorio_cnpj = adv.get("escritorio_cnpj", "")
    if escritorio_nome:
        texto += (
            f", neste ato representando o escritório {escritorio_nome}, "
            f"pessoa jurídica de direito privado, inscrito no CNPJ sob o nº {escritorio_cnpj}"
        )
    return texto


def _qualificar_advogado_marcado(adv: dict) -> str:
    """Igual a _qualificar_advogado, mas com marcadores de negrito no nome, OAB,
    escritório e CNPJ (aplicados no render por aplicar_marcadores_negrito)."""
    genero = adv.get("genero", "masculino")
    inscrito = "inscrita" if genero == "feminino" else "inscrito"
    texto = (
        f"{marcar_negrito(adv.get('nome_completo', ''))}, "
        f"{adv.get('nacionalidade', '')}, "
        f"{adv.get('estado_civil', '')}, "
        f"advogado, {inscrito} na {marcar_negrito(adv.get('numero_oab', ''))}"
    )
    escritorio_nome = adv.get("escritorio_nome", "")
    escritorio_cnpj = adv.get("escritorio_cnpj", "")
    if escritorio_nome:
        texto += (
            f", neste ato representando o escritório {marcar_negrito(escritorio_nome)}, "
            f"pessoa jurídica de direito privado, inscrito no CNPJ sob o nº {marcar_negrito(escritorio_cnpj)}"
        )
    return texto


def _build_telefone_text(c, is_masc: bool) -> str:
    """Constrói o texto de telefone no mesmo formato do template de kit."""
    telefones: list[dict] = []
    if c.telefone:
        telefones.append({
            "numero": c.telefone,
            "titular_contato": c.titular_contato or "sim",
            "nome_titular_numero": c.nome_titular_numero or "",
            "relacao_titular_tipo": c.relacao_titular_tipo or "",
            "relacao_titular": c.relacao_titular or "",
        })
    for t in (c.telefones_extras or []):
        if (t.get("numero") or "").strip():
            telefones.append(t)

    validos = [t for t in telefones if (t.get("numero") or "").strip()]
    if not validos:
        return ""

    terceiros = [
        t for t in validos
        if t.get("titular_contato") == "nao"
        and (t.get("nome_titular_numero") or "").strip()
        and _formatar_relacao(
            t.get("relacao_titular_tipo", ""),
            t.get("relacao_titular", ""),
        )
    ]
    proprios = [t for t in validos if t not in terceiros]

    partes: list[str] = []
    if proprios:
        nums = [t["numero"] for t in proprios]
        sufixo = "pertencente" if len(nums) == 1 else "pertencentes"
        partes.append(f"{_juntar(nums)}, {sufixo} ao cliente")
    for t in terceiros:
        relacao = _formatar_relacao(
            t.get("relacao_titular_tipo", ""),
            t.get("relacao_titular", ""),
        )
        gen_suf = "do" if is_masc else "da"
        partes.append(
            f"{t['numero']}, pertencente à "
            f"{(t.get('nome_titular_numero') or '').strip()}, "
            f"{relacao} {gen_suf} cliente"
        )
    return _juntar(partes)


# ---------------------------------------------------------------------------
# Construção do contexto Jinja2
# ---------------------------------------------------------------------------

def build_kit_context(kit: Kit) -> dict:
    """Porta montarContexto() do NovoKitView.vue para Python."""
    c = kit.cliente
    is_masc = (c.genero or "masculino") == "masculino"
    acoes = list(kit.acoes.prefetch_related("associacao").all())
    snapshot: list[dict] = kit.advogados_snapshot or []
    uf = (c.uf or "").upper()
    cidade = c.cidade or ""
    today = date.today()

    # Flexão de gênero
    inscrito = "inscrito" if is_masc else "inscrita"
    domiciliado = "domiciliado" if is_masc else "domiciliada"
    contratado = "contratado" if is_masc else "contratada"

    # Endereço
    numero = c.numero or ""
    if numero and re.match(r"^s\s*/?\s*n$", numero, re.IGNORECASE):
        numero_fmt = numero
    elif numero:
        numero_fmt = f"nº {numero}"
    else:
        numero_fmt = ""
    partes_end = [p for p in [c.logradouro, numero_fmt, c.complemento, c.bairro] if p]
    endereco = f"{', '.join(partes_end)}, {cidade}/{uf}, CEP {_fmt_cep(c.cep or '')}"

    # Data/local
    local_data = f"{cidade}/{uf}, {today.day} de {MESES[today.month - 1]} de {today.year}"

    # Bancos das ações (deduplica, preserva ordem)
    banco_nomes: list[str] = []
    for a in acoes:
        if a.tipo_acao == "contribuicao_sindical_nao_autorizada" and kit.tipo == "bancario" and a.associacao_id:
            nome = (a.associacao.abreviacao or a.associacao.nome).upper().strip()
        else:
            nome = a.banco_outro if a.nome_banco == "Outro" else a.nome_banco
            nome = (nome or "").upper().strip()
        banco_nomes.append(nome)
    seen: set[str] = set()
    banco_nomes_uniq: list[str] = []
    for n in banco_nomes:
        if n and n not in seen:
            seen.add(n)
            banco_nomes_uniq.append(n)

    cliente_am = uf == "AM"
    if not banco_nomes_uniq:
        bancos = "(BANCOS DA AÇÃO)"
    elif kit.tipo == "bancario" and not cliente_am:
        if len(banco_nomes_uniq) == 1:
            bancos = f"{banco_nomes_uniq[0]} e INSS – Instituto Nacional do Seguro Social"
        else:
            bancos = ", ".join(banco_nomes_uniq) + " e INSS – Instituto Nacional do Seguro Social"
    else:
        bancos = _juntar(banco_nomes_uniq)

    # Tipos de ação
    tipo_labels = [TIPOS_ACAO_LABELS.get(a.tipo_acao, a.tipo_acao) for a in acoes]
    tipos_acao = _juntar(tipo_labels)

    # Tarifas questionadas
    tarifa_labels: list[str] = []
    for a in acoes:
        if a.tipo_acao == "tarifa_bancaria" and a.tarifa_questionada:
            t = a.tarifa_questionada_outro if a.tarifa_questionada == "OUTROS" else a.tarifa_questionada
            if t:
                tarifa_labels.append(t)
    tarifas_questionadas = _juntar(tarifa_labels)
    tarifa_questionada = tarifa_labels[0] if tarifa_labels else ""

    # Advogados (snapshot já armazenado no kit)
    socios = sorted(
        [a for a in snapshot if a.get("is_socio")],
        key=lambda x: x.get("nome_completo", ""),
    )
    nao_socios = [a for a in snapshot if not a.get("is_socio")]

    _placeholder = "(OAB REFERENTE AO ESTADO DA AÇÃO)"
    oab_estado = _placeholder
    contratados_socios = ""
    advogados_estado = ""
    unidade_apoio = ""
    socio1_oab = _placeholder
    socio2_oab = _placeholder
    oab_tiago = _placeholder
    oab_eduardo = _placeholder

    if socios:
        oab_estado = " e ".join(s.get("numero_oab", "") for s in socios)
        contratados_socios = " e ".join(_qualificar_advogado(s) for s in socios)
        socio1_oab = socios[0].get("numero_oab", _placeholder)
        socio2_oab = socios[1].get("numero_oab", _placeholder) if len(socios) > 1 else _placeholder
        tiago = next(
            (s for s in socios if "tiago" in _normalize(s.get("nome_completo", ""))), None
        )
        eduardo = next(
            (s for s in socios if "eduardo" in _normalize(s.get("nome_completo", ""))), None
        )
        if tiago:
            oab_tiago = tiago.get("numero_oab", _placeholder)
        if eduardo:
            oab_eduardo = eduardo.get("numero_oab", _placeholder)

    if nao_socios:
        advogados_estado = "; e ".join(_qualificar_advogado_marcado(a) for a in nao_socios)

    # A unidade de apoio administrativo deve corresponder ao estado da ação
    # (UF do cliente), não a uma OAB resolvida por fallback (SC ou 1ª cadastrada).
    # Sem esse guard, um advogado sem OAB na UF do cliente injetaria a unidade
    # de outro estado (ex.: Arapiraca/AL num kit de cliente de SC).
    for adv in snapshot:
        if adv.get("oab_fonte") == "uf_cliente" and adv.get("unidade_apoio_endereco"):
            unidade_apoio = f", unidade de apoio administrativo na {adv['unidade_apoio_endereco']}"
            break

    # Bloco de assinatura
    condicao = c.condicao_cliente or "alfabetizado"
    bloco_assinatura = "_________________________________________________\nASSINANTE"
    if condicao == "analfabeto":
        t1n = c.testemunha1_nome or ""
        t1c = _fmt_cpf(c.testemunha1_cpf or "")
        t2n = c.testemunha2_nome or ""
        t2c = _fmt_cpf(c.testemunha2_cpf or "")
        bloco_assinatura = (
            "QUANDO ANALFABETO:\n\n"
            "_________________________________________________\nDeclarante/Rogado\n\n"
            f"_________________________________________________\nTestemunha 1 Nome/CPF: {t1n} / {t1c}\n\n"
            f"_________________________________________________\nTestemunha 2 Nome/CPF: {t2n} / {t2c}"
        )
    elif condicao in ("incapaz", "crianca_adolescente"):
        label = (
            "Responsável Legal (Criança/Adolescente)"
            if condicao == "crianca_adolescente"
            else "Responsável Legal (Incapaz)"
        )
        bloco_assinatura = (
            f"_________________________________________________\n"
            f"{label} - {c.responsavel_legal_nome or ''} - CPF: {_fmt_cpf(c.responsavel_legal_cpf or '')}"
        )

    # Bloco de assinatura do domicílio
    bloco_assinatura_domicilio = bloco_assinatura
    if condicao == "analfabeto":
        bloco_assinatura_domicilio = (
            "QUANDO ANALFABETO:\n\n"
            "__________________________________________________\nDeclarante/titular do comprovante de endereço\n\n"
            "__________________________________________________\nAssinatura do rogado\n\n"
            f"TESTEMUNHA: {c.testemunha1_nome or ''} CPF: {_fmt_cpf(c.testemunha1_cpf or '')}\n"
            f"TESTEMUNHA: {c.testemunha2_nome or ''} CPF: {_fmt_cpf(c.testemunha2_cpf or '')}"
        )

    # Hipossuficiência
    def _sim_nao(val) -> str:
        return "( X ) SIM    (   ) NÃO" if val else "(   ) SIM    ( X ) NÃO"

    # Data de nascimento
    data_nasc = c.data_nascimento.strftime("%d/%m/%Y") if c.data_nascimento else ""

    return {
        "nome_cliente": c.nome_completo.upper(),
        "nacionalidade": c.nacionalidade or "",
        "data_nascimento": data_nasc,
        "estado_civil": c.estado_civil or "",
        "profissao": c.profissao or "",
        "cpf_cliente": _fmt_cpf(c.cpf or ""),
        "endereco": endereco,
        "telefone": _build_telefone_text(c, is_masc),
        "inscrito": inscrito,
        "domiciliado": domiciliado,
        "contratado": contratado,
        "o_contratante": "O" if is_masc else "A",
        "pelo": "pelo" if is_masc else "pela",
        "ao": "ao" if is_masc else "à",
        "do": "do" if is_masc else "da",
        "no": "no" if is_masc else "na",
        "informado": "informado" if is_masc else "informada",
        "impossibilitado": "impossibilitado" if is_masc else "impossibilitada",
        "oab_estado": oab_estado,
        "socio1_oab": socio1_oab,
        "socio2_oab": socio2_oab,
        "oab_tiago": oab_tiago,
        "oab_eduardo": oab_eduardo,
        "unidade_apoio": unidade_apoio,
        "contratados_socios": contratados_socios,
        "advogados_estado": advogados_estado,
        "bancos": bancos,
        "tipos_acao": tipos_acao,
        "tarifa_questionada": tarifa_questionada,
        "tarifas_questionadas": tarifas_questionadas,
        "instituicao_re": bancos,
        "imoveis_sim_nao": _sim_nao(c.possui_imoveis),
        "moveis_sim_nao": _sim_nao(c.possui_moveis),
        "irpf_sim_nao": _sim_nao(c.isento_irpf),
        "local_data": local_data,
        "bloco_assinatura": bloco_assinatura,
        "responsavel_imovel_nome": c.responsavel_imovel_nome or "",
        "responsavel_imovel_cpf": _fmt_cpf(c.responsavel_imovel_cpf or ""),
        "bloco_assinatura_domicilio": bloco_assinatura_domicilio,
        "condicao_cliente": condicao,
        "rogado_nome": c.rogado_nome or "",
        "rogado_cpf": _fmt_cpf(c.rogado_cpf or ""),
        "testemunha1_nome": c.testemunha1_nome or "",
        "testemunha1_cpf": _fmt_cpf(c.testemunha1_cpf or ""),
        "testemunha2_nome": c.testemunha2_nome or "",
        "testemunha2_cpf": _fmt_cpf(c.testemunha2_cpf or ""),
        "responsavel_legal_nome": c.responsavel_legal_nome or "",
        "responsavel_legal_cpf": _fmt_cpf(c.responsavel_legal_cpf or ""),
        "clausula_porcentagem": kit.clausula_porcentagem_snapshot or "",
        "honorarios_iniciais": _texto_honorarios_iniciais(kit),
        # Palavra seguinte no template: minúscula quando há cláusula de honorários
        # (o texto termina em vírgula), maiúscula quando começa frase nova.
        "honorarios_pelos": "pelos" if _tem_honorarios_iniciais(kit) else "Pelos",
        # Defaults de procuração (sobrescritos pelo build_procuracao_context por ação)
        "procuracao_frase_acao": "",
        "procuracao_usa_numero_contrato": False,
        "procuracao_detalhe_tipo": "",
        "contrato_acao": "",
    }


# ---------------------------------------------------------------------------
# Notificação extrajudicial (template_notificacao_extrajudicial)
# ---------------------------------------------------------------------------

# Mapa tipo_acao -> bloco {%p if tipo_notificacao == '...' %} do template.
# TODO: definir os pendentes (cartao_credito_consignado,
# contribuicao_sindical_nao_autorizada, seguro_nao_autorizado) e os blocos
# ainda sem par (encargos, pagamento_eletronico, titulo_capitalizacao).
TIPO_ACAO_NOTIFICACAO = {
    "rmc": "rmc",
    "rcc": "rcc",
    "tarifa_bancaria": "tarifas",
    "emprestimo_nao_autorizado": "nao_contratado",
}


def _adv_notificacao(snapshot: list[dict]) -> dict:
    """Advogado que assina a notificação: OAB na UF do cliente, senão o 1º."""
    if not snapshot:
        return {}
    for a in snapshot:
        if a.get("oab_fonte") == "uf_cliente":
            return a
    return snapshot[0]


def _fmt_moeda(v: int) -> str:
    """15000 -> 'R$ 15.000,00'."""
    s = f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"R$ {s}"


# Extenso (pt-BR) para os valores de danos morais — sempre múltiplos de mil.
_EXT_U = ["", "um", "dois", "três", "quatro", "cinco", "seis", "sete", "oito",
          "nove", "dez", "onze", "doze", "treze", "quatorze", "quinze",
          "dezesseis", "dezessete", "dezoito", "dezenove"]
_EXT_D = ["", "", "vinte", "trinta", "quarenta", "cinquenta", "sessenta",
          "setenta", "oitenta", "noventa"]
_EXT_C = ["", "cento", "duzentos", "trezentos", "quatrocentos", "quinhentos",
          "seiscentos", "setecentos", "oitocentos", "novecentos"]


def _extenso_ate_999(n: int) -> str:
    if n == 0:
        return ""
    if n == 100:
        return "cem"
    partes = []
    c, d = divmod(n, 100)
    if c:
        partes.append(_EXT_C[c])
    if d:
        if d < 20:
            partes.append(_EXT_U[d])
        else:
            dez, uni = divmod(d, 10)
            partes.append(f"{_EXT_D[dez]} e {_EXT_U[uni]}" if uni else _EXT_D[dez])
    return " e ".join(partes)


def _extenso_reais_mil(v: int) -> str:
    """Extenso de valor múltiplo de mil: 25000 -> 'vinte e cinco mil reais'."""
    milhares = v // 1000
    return "mil reais" if milhares == 1 else f"{_extenso_ate_999(milhares)} mil reais"


# ── Moeda e extenso com centavos (espelha useCurrencyMask/useNumeroExtenso do front) ──

def _num_milhares(n: int) -> str:
    """Inteiro por extenso (mesma lógica de converterMilhares do front)."""
    if n < 1000:
        return _extenso_ate_999(n)
    if n < 1_000_000:
        mil, resto = divmod(n, 1000)
        mil_str = "um mil" if mil == 1 else f"{_extenso_ate_999(mil)} mil"
        return mil_str + (f" {_extenso_ate_999(resto)}" if resto else "")
    if n < 1_000_000_000:
        milhao, resto = divmod(n, 1_000_000)
        milhao_str = "um milhão" if milhao == 1 else f"{_extenso_ate_999(milhao)} milhões"
        return milhao_str + (f" {_num_milhares(resto)}" if resto else "")
    return ""


def _moeda_br(valor: Decimal) -> str:
    """Decimal -> '1.234,56' (mesma saída do formatCurrency do front)."""
    v = Decimal(valor).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    inteiro = int(v)
    centavos = int((v - inteiro) * 100)
    return f"{inteiro:,}".replace(",", ".") + f",{centavos:02d}"


def _extenso_moeda(valor: Decimal) -> str:
    """Decimal -> extenso com reais e centavos (espelha numeroParaExtenso do front)."""
    v = Decimal(valor).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    inteiro = int(v)
    centavos = int((v - inteiro) * 100)
    inteiro_str = _num_milhares(inteiro)
    centavos_str = _extenso_ate_999(centavos) if centavos > 0 else ""
    res = ""
    if inteiro_str:
        res += inteiro_str + (" real" if inteiro == 1 else " reais")
    if centavos_str:
        if res:
            res += " e "
        res += centavos_str + (" centavo" if centavos == 1 else " centavos")
    return res or "zero"


_HONORARIOS_INICIAIS_TEXTO = (
    "O(A) CONTRATANTE pagará honorários iniciais destinados à análise técnica, "
    "levantamento bancário, organização documental e encaminhamento das medidas "
    "judiciais e/ou extrajudiciais cabíveis relacionadas aos contratos/empréstimos "
    "não reconhecidos, no importe de R$ {valor} ({extenso}). Ademais, O(A) CONTRATANTE "
    "autoriza que referido valor seja compensado/descontado de eventual crédito, "
    "prestação de contas, acordo ou alvará existente em seu favor junto ao CONTRATADO,"
)


def _tem_honorarios_iniciais(kit: Kit) -> bool:
    valor = getattr(kit, "honorarios_iniciais", None)
    return bool(valor) and Decimal(valor) > 0


def _texto_honorarios_iniciais(kit: Kit) -> str:
    """Cláusula de honorários iniciais para o contrato. Vazio se não houver valor."""
    if not _tem_honorarios_iniciais(kit):
        return ""
    valor = kit.honorarios_iniciais
    return _HONORARIOS_INICIAIS_TEXTO.format(
        valor=_moeda_br(valor), extenso=_extenso_moeda(valor),
    )


def _banco_acao(a) -> str:
    """Nome do banco da ação (normalizado), considerando 'Outro'."""
    return _normalize(a.banco_outro if a.nome_banco == "Outro" else a.nome_banco)


def _valor_danos_morais_doc(kit: Kit, acoes: list) -> int:
    """Danos morais do DOCUMENTO, em reais (0 se não se aplica).

    - RCC / RMC: 15000.
    - nao_contratado: o documento agrupa as ações do mesmo banco → 5000 × nº de
      ações quando há mais de uma; senão 15000.
    - tarifa_bancaria (um doc por ação): 5000 se o banco se repete entre as
      ações de tarifa do kit; senão 15000.
    """
    rep = acoes[0]
    tipo = rep.tipo_acao
    if tipo in ("rcc", "rmc"):
        return 15000
    if tipo == "emprestimo_nao_autorizado":
        return 5000 * len(acoes) if len(acoes) > 1 else 15000
    if tipo == "tarifa_bancaria":
        alvo = _banco_acao(rep)
        mesmo_tipo = [a for a in kit.acoes.all() if a.tipo_acao == tipo]
        repetido = sum(1 for a in mesmo_tipo if _banco_acao(a) == alvo) > 1
        return 5000 if repetido else 15000
    return 0


def build_notificacao_context(kit: Kit, acoes: list) -> dict:
    """Contexto (placeholders {{MAIÚSCULO}}) do template de notificação, por documento.

    `acoes` é a lista de ações que compõem UM documento: para nao_contratado do
    mesmo banco, várias ações (contratos) juntas; para os demais tipos, uma só.

    Origens (definidas com o operador):
      - Cliente: dados pessoais + endereço + assinatura (cidade/UF do cliente)
      - Advogado: snapshot do kit
      - Banco: BancoKit (casado por nome_banco da ação)
      - NUMERO_CONTRATO: todos os contratos das ações do documento
    """
    from .models import BancoKit

    c = kit.cliente
    adv = _adv_notificacao(kit.advogados_snapshot or [])
    today = date.today()
    rep = acoes[0]  # ação representativa (banco/tipo iguais no grupo)

    nome_banco = rep.banco_outro if rep.nome_banco == "Outro" else rep.nome_banco
    banco = BancoKit.objects.filter(nome=nome_banco).first() if nome_banco else None

    partes_end = [c.logradouro or ""]
    if c.numero:
        partes_end.append(f"nº {c.numero}")
    if c.complemento:
        partes_end.append(c.complemento)
    endereco_cliente = ", ".join(p for p in partes_end if p).strip(", ")

    nome_adv = adv.get("nome_completo", "")
    _danos = _valor_danos_morais_doc(kit, acoes)
    numeros = ", ".join(a.numero_contrato for a in acoes if a.numero_contrato)

    return {
        "tipo_notificacao": TIPO_ACAO_NOTIFICACAO.get(rep.tipo_acao, ""),
        # Cliente
        "NOME_CLIENTE": c.nome_completo or "",
        "CPF": _fmt_cpf(c.cpf or ""),
        "RG": c.rg or "",
        "ORGAO_EXPEDIDOR": c.orgao_expedidor or "",
        "NACIONALIDADE": c.nacionalidade or "",
        "ESTADO_CIVIL": c.estado_civil or "",
        "PROFISSAO": c.profissao or "",
        "ENDERECO_CLIENTE": endereco_cliente,
        "BAIRRO": c.bairro or "",
        "MUNICIPIO": c.cidade or "",
        "CEP_CLIENTE": _fmt_cep(c.cep or ""),
        "UF_CLIENTE": (c.uf or "").upper(),
        # Assinatura = cidade/UF do cliente
        "CIDADE_ASSINATURA": c.cidade or "",
        "UF_ASSINATURA": (c.uf or "").upper(),
        "DATA_EXTENSO": f"{today.day:02d} de {MESES[today.month - 1]} de {today.year}",
        # Advogado (snapshot)
        "NOME_ADVOGADO": nome_adv,
        "NOME_ADVOGADO_MAIUSCULO": nome_adv.upper(),
        "OAB_ADVOGADO": adv.get("numero_oab", ""),
        "ENDERECO_ESCRITORIO": adv.get("escritorio_endereco", ""),
        # Banco (BancoKit)
        "NOME_BANCO": (banco.nome if banco else nome_banco) or "",
        "NOME_BANCO_CONTRATO": (banco.nome if banco else nome_banco) or "",
        "CNPJ_BANCO": banco.cnpj if banco else "",
        "ENDERECO_BANCO_LINHA1": banco.endereco if banco else "",
        "ENDERECO_BANCO_LINHA2": "",
        "CEP_BANCO": "",
        "CIDADE_BANCO": "",
        "UF_BANCO": "",
        # Ação (todos os contratos do documento)
        "NUMERO_CONTRATO": numeros,
        # Empréstimo — pendente (a definir)
        "VALOR_EMPRESTIMO": "",
        "QTD_PARCELAS": "",
        "VALOR_PARCELA": "",
        "DATA_INCLUSAO": "",
        "DATA_INICIO_DESCONTOS": "",
        # Danos morais (regra por documento — soma no nao_contratado agrupado)
        "VALOR_DANOS_MORAIS": _fmt_moeda(_danos) if _danos else "",
        "VALOR_DANOS_MORAIS_EXTENSO": _extenso_reais_mil(_danos) if _danos else "",
    }


def _find_notificacao_template():
    """Template ativo cujo nome contém 'notificacao extrajudicial'."""
    for tpl in Template.objects.filter(active=True):
        n = _normalize(tpl.name)
        if all(w in n for w in ("notificacao", "extrajudicial")):
            return tpl
    return None


def _merge_pdfs(pdf_list: list[bytes]) -> bytes:
    """Concatena PDFs já renderizados (uma notificação por ação)."""
    from pypdf import PdfReader, PdfWriter
    writer = PdfWriter()
    for pdf in pdf_list:
        for page in PdfReader(BytesIO(pdf)).pages:
            writer.add_page(page)
    out = BytesIO()
    writer.write(out)
    return out.getvalue()


def notificacao_documentos(kit: Kit) -> list[dict]:
    """Lista os DOCUMENTOS de notificação do kit.

    Cada item: {"rep_id": <id da 1ª ação>, "tipo_acao", "banco", "acoes": [...]}.
    - emprestimo_nao_autorizado: ações do MESMO banco viram um único documento
      (todos os contratos juntos).
    - demais tipos mapeados: um documento por ação.
    """
    mapeadas = [a for a in kit.acoes.all() if TIPO_ACAO_NOTIFICACAO.get(a.tipo_acao)]

    docs: list[dict] = []
    grupos_nc: dict[str, list] = {}
    for a in mapeadas:
        banco_label = (a.banco_outro if a.nome_banco == "Outro" else a.nome_banco) or ""
        if a.tipo_acao == "emprestimo_nao_autorizado":
            grupos_nc.setdefault(_banco_acao(a), []).append(a)
        else:
            docs.append({
                "rep_id": a.id, "tipo_acao": a.tipo_acao,
                "banco": banco_label, "acoes": [a],
            })
    for lst in grupos_nc.values():
        rep = lst[0]
        banco_label = (rep.banco_outro if rep.nome_banco == "Outro" else rep.nome_banco) or ""
        docs.append({
            "rep_id": rep.id, "tipo_acao": rep.tipo_acao,
            "banco": banco_label, "acoes": lst,
        })
    docs.sort(key=lambda d: d["rep_id"])
    return docs


def gerar_notificacao_pdf(kit: Kit, acao=None) -> bytes:
    """Gera o PDF da notificação extrajudicial.

    - acao=None: todos os documentos do kit, combinados num PDF único.
    - acao=<AcaoKit>: só o documento que contém essa ação (nao_contratado do
      mesmo banco é um documento só).

    Levanta FileNotFoundError se o template não existir e ValueError se não
    houver ação com tipo de notificação mapeado.
    """
    tpl = _find_notificacao_template()
    if tpl is None:
        raise FileNotFoundError(
            "Template 'template_notificacao_extrajudicial' não encontrado ou inativo."
        )
    docs = notificacao_documentos(kit)
    if acao is not None:
        docs = [d for d in docs if any(a.id == acao.id for a in d["acoes"])]
    if not docs:
        raise ValueError(
            "Nenhuma ação deste kit tem tipo de notificação definido ainda."
        )
    pdfs = []
    for d in docs:
        ctx = build_notificacao_context(kit, d["acoes"])
        docx_bytes = _render_template_to_docx(tpl, ctx)
        pdfs.append(_docx_to_pdf(docx_bytes))
    return pdfs[0] if len(pdfs) == 1 else _merge_pdfs(pdfs)


# ---------------------------------------------------------------------------
# Contexto por ação — procuração
# ---------------------------------------------------------------------------

def _detalhe_bruto_procuracao(acao) -> str:
    """Retorna o texto bruto de detalhe da ação para uso na procuração."""
    if acao.tipo_acao == "tarifa_bancaria":
        if acao.tarifa_questionada == "OUTROS":
            return (acao.tarifa_questionada_outro or "").strip()
        return acao.tarifa_questionada or ""
    if acao.tipo_acao == "seguro_nao_autorizado":
        return (acao.tipo_seguro or "").strip()
    if acao.tipo_acao == "contribuicao_sindical_nao_autorizada":
        return (acao.tipo_contribuicao or "").strip()
    return ""


def build_procuracao_context(base_context: dict, acao, kit: Kit) -> dict:
    """Porta montarContextoProcuracao() do frontend — contexto por ação para a procuração."""
    banco_nome = (acao.banco_outro if acao.nome_banco == "Outro" else acao.nome_banco or "").upper()
    tipo_label = TIPOS_ACAO_LABELS.get(acao.tipo_acao, acao.tipo_acao)
    usa_contrato = acao.tipo_acao in TIPOS_COM_CONTRATO
    procuracao_detalhe_tipo = "" if usa_contrato else _detalhe_bruto_procuracao(acao).strip()
    banco_com_inss = banco_nome  # procuração nunca leva "e INSS"
    do_pronome = base_context.get("do", "do")

    # Caso especial: contribuição sindical com associação em kit bancário
    if (
        acao.tipo_acao == "contribuicao_sindical_nao_autorizada"
        and kit.tipo == "bancario"
        and acao.associacao_id
    ):
        tipo_label = "contribuição não autorizada em benefício previdenciário"
        banco_com_inss = (acao.associacao.abreviacao or acao.associacao.nome).upper().strip()
        procuracao_detalhe_tipo = ""
        do_pronome = "de(a)"

    tarifa = ""
    if acao.tipo_acao == "tarifa_bancaria" and acao.tarifa_questionada:
        tarifa = (
            (acao.tarifa_questionada_outro or "").strip()
            if acao.tarifa_questionada == "OUTROS"
            else acao.tarifa_questionada
        )

    # frase quando não há número de contrato
    procuracao_frase_acao = (
        ""
        if usa_contrato
        else (f" de {procuracao_detalhe_tipo}" if procuracao_detalhe_tipo else "")
    )

    return {
        **base_context,
        "do": do_pronome,
        "bancos": banco_com_inss,
        "tipos_acao": tipo_label,
        "tarifa_questionada": tarifa,
        "instituicao_re": banco_com_inss,
        "contrato_acao": f"contrato nº {acao.numero_contrato or ''}",
        "procuracao_usa_numero_contrato": usa_contrato,
        "procuracao_detalhe_tipo": procuracao_detalhe_tipo,
        "procuracao_frase_acao": procuracao_frase_acao,
    }


# ---------------------------------------------------------------------------
# Resolução de templates
# ---------------------------------------------------------------------------

def find_kit_templates(tipo: str) -> dict[str, Template]:
    """Retorna {key: Template} para o tipo de kit, usando keyword matching."""
    defs = KIT_TEMPLATE_DEFS.get(tipo, KIT_TEMPLATE_DEFS["bancario"])
    all_templates = list(Template.objects.filter(active=True))
    result: dict[str, Template] = {}

    for defn in defs:
        keyword_words = _normalize(defn["keyword"]).split()
        for tpl in all_templates:
            name = _normalize(tpl.name)
            if not all(w in name for w in keyword_words):
                continue
            # Exclui templates de outros tipos (mesma regra do frontend)
            if tipo == "bancario" and ("previdenciario" in name or "marketing" in name):
                continue
            if tipo == "previdenciario" and "marketing" in name:
                continue
            result[defn["key"]] = tpl
            break  # primeiro match vence

    return result


# ---------------------------------------------------------------------------
# Renderização
# ---------------------------------------------------------------------------

def _render_template_to_docx(tpl: Template, context: dict) -> bytes:
    if DocxTemplate is None:
        raise RuntimeError("docxtpl não instalado no servidor.")

    file_path = Path(tpl.file.path)
    if not file_path.exists():
        raise FileNotFoundError(f"Arquivo do template '{tpl.name}' não encontrado no servidor.")

    normalized_path = None
    try:
        normalized_path = normalize_docx_jinja_runs(file_path)
        doc = DocxTemplate(str(normalized_path))
        env = build_env()
        doc.render(context, jinja_env=env)
        aplicar_marcadores_negrito(doc.docx)
        buf = BytesIO()
        doc.save(buf)
        buf.seek(0)
        try:
            buf = strip_blank_pages(buf)
        except Exception:
            buf.seek(0)
        return buf.read()
    finally:
        if normalized_path is not None:
            normalized_path.unlink(missing_ok=True)


def _compose_docx_files(docx_bytes_list: list[bytes]) -> bytes:
    """Combina múltiplos .docx em um único documento com quebra de página entre eles."""
    import copy
    from docx import Document
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    master = Document(BytesIO(docx_bytes_list[0]))
    master_body = master.element.body
    final_sect_pr = master_body.find(qn("w:sectPr"))

    def add_to_master(element):
        if final_sect_pr is not None:
            final_sect_pr.addprevious(element)
        else:
            master_body.append(element)

    for raw in docx_bytes_list[1:]:
        page_break = OxmlElement("w:p")
        run_elem = OxmlElement("w:r")
        br = OxmlElement("w:br")
        br.set(qn("w:type"), "page")
        run_elem.append(br)
        page_break.append(run_elem)
        add_to_master(page_break)

        doc = Document(BytesIO(raw))
        for child in doc.element.body:
            tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag
            if tag in ("p", "tbl"):
                add_to_master(copy.deepcopy(child))

    buf = BytesIO()
    master.save(buf)
    return buf.getvalue()


def _run_libreoffice(docx_bytes: bytes) -> bytes:
    """Converte .docx → PDF via LibreOffice headless. Não altera o .docx."""
    with tempfile.TemporaryDirectory() as tmpdir:
        input_path = Path(tmpdir) / "input.docx"
        input_path.write_bytes(docx_bytes)

        profile_dir = Path(tmpdir) / "profile"
        profile_dir.mkdir()
        user_install = f"-env:UserInstallation=file://{profile_dir}"

        result = subprocess.run(
            [
                "libreoffice",
                user_install,
                "--headless",
                "--norestore",
                "--nolockcheck",
                "--nodefault",
                "--convert-to", "pdf",
                "--outdir", tmpdir,
                str(input_path),
            ],
            capture_output=True,
            timeout=120,
            check=False,
        )

        pdf_path = Path(tmpdir) / "input.pdf"
        if not pdf_path.exists():
            stderr = (result.stderr or b"").decode("utf-8", errors="replace").strip()
            stdout = (result.stdout or b"").decode("utf-8", errors="replace").strip()
            detail = stderr or stdout or "saída vazia"
            raise RuntimeError(
                f"Falha na conversão para PDF (exit={result.returncode}): {detail[:500]}"
            )
        return pdf_path.read_bytes()


def _docx_to_pdf(docx_bytes: bytes) -> bytes:
    """Converte .docx → PDF com numeração de páginas condicional.

    Injeta 'Página X de Y' no rodapé somente quando o documento tem mais de
    uma página — mesma estratégia de _convert_with_conditional_page_numbering
    em templates_app/views.py. Documentos de página única não recebem numeração.
    """
    from pypdf import PdfReader
    from templates_app.docx_page_numbering import add_page_numbering_simple

    docx_bytes = flatten_inherited_formatting(docx_bytes)
    pdf_bytes = _run_libreoffice(docx_bytes)

    try:
        page_count = len(PdfReader(BytesIO(pdf_bytes)).pages)
    except Exception:
        return pdf_bytes

    if page_count <= 1:
        return pdf_bytes

    return _run_libreoffice(add_page_numbering_simple(docx_bytes))


# ---------------------------------------------------------------------------
# Orquestrador principal
# ---------------------------------------------------------------------------

def gerar_documentos_kit(kit: Kit) -> tuple[list[DocumentoKit], list[str]]:
    """
    Gera todos os PDFs do kit e salva como DocumentoKit.

    - Apaga os DocumentoKit anteriores antes de gerar.
    - domicilio: gerado só se cliente.comprovante_nome_cliente == 'nao'.
    - Retorna (documentos_criados, warnings).
    """
    if DocxTemplate is None:
        raise RuntimeError("docxtpl não instalado no servidor.")

    # Validação de pré-condições por tipo de kit
    if kit.tipo == "previdenciario" and not kit.cliente.data_nascimento:
        raise RuntimeError(
            "Kits previdenciários exigem a data de nascimento do cliente preenchida. "
            "Cadastre a data de nascimento antes de gerar os documentos."
        )

    templates = find_kit_templates(kit.tipo)
    if not templates:
        raise RuntimeError(
            f"Nenhum template encontrado para kit do tipo '{kit.tipo}'. "
            "Cadastre os templates de kit no JurisDoc e marque-os como ativos."
        )

    context = build_kit_context(kit)
    acoes = list(kit.acoes.prefetch_related("associacao").all())
    include_domicilio = kit.cliente.comprovante_nome_cliente == "nao"
    valid_tipos = {v for v, _ in DocumentoKit.TIPO_CHOICES}

    # Remove documentos anteriores — apaga os arquivos físicos antes dos registros,
    # senão o Django adiciona sufixo aleatório ao salvar um arquivo com mesmo nome.
    for _doc in kit.documentos.all():
        _doc.arquivo.delete(save=False)
    kit.documentos.all().delete()

    created: list[DocumentoKit] = []
    warnings: list[str] = []

    for key, tpl in templates.items():
        if key == "domicilio" and not include_domicilio:
            continue
        if key not in valid_tipos:
            warnings.append(f"Tipo '{key}' não reconhecido nos choices de DocumentoKit — ignorado.")
            continue

        try:
            # Procuração bancária/marketing: uma página por ação, combinadas em um único PDF
            if key == "procuracao" and kit.tipo != "previdenciario" and acoes:
                action_docx_list: list[bytes] = []
                for acao in acoes:
                    per_ctx = build_procuracao_context(context, acao, kit)
                    action_docx_list.append(_render_template_to_docx(tpl, per_ctx))
                final_docx = (
                    action_docx_list[0]
                    if len(action_docx_list) == 1
                    else _compose_docx_files(action_docx_list)
                )
                pdf_bytes = _docx_to_pdf(final_docx)
            else:
                docx_bytes = _render_template_to_docx(tpl, context)
                pdf_bytes = _docx_to_pdf(docx_bytes)
        except FileNotFoundError as e:
            warnings.append(str(e))
            continue
        except RuntimeError as e:
            warnings.append(f"Erro ao gerar '{key}': {e}")
            continue
        except Exception as e:
            warnings.append(f"Erro inesperado ao gerar '{key}': {e}")
            continue

        cliente_slug = slug_nome_cliente(kit.cliente.nome_completo or "")
        doc = DocumentoKit(kit=kit, tipo=key)
        doc.arquivo.save(
            f"kits/{kit.id}/{key}_{cliente_slug}.pdf",
            ContentFile(pdf_bytes),
            save=True,
        )
        created.append(doc)

    return created, warnings
